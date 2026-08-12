"""
Excel export — a summary sheet plus one detail sheet per project.

Durations are written as real numbers (decimal hours) rather than as text, so
the recipient can sum, pivot and invoice from them without retyping anything.
"""
import re
from typing import Optional

from ..database import Database
from ..models import Project, format_hm
from ..timeutil import fmt_date, fmt_time, now_local, parse_local
from .common import period_label

# Excel forbids these in a sheet name, and caps the name at 31 characters.
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
MAX_SHEET_NAME = 31

# ── Palette (Excel wants AARRGGBB) ───────────────────────────────────────────
C_BAND = "FF0F2A47"
C_HEAD = "FF39424F"
C_WHITE = "FFFFFFFF"
C_INK = "FF12161C"
C_MUTED = "FF6B7684"
C_ACCENT = "FF2C7BE5"
C_ZEBRA = "FFF7F9FC"


def _safe_sheet_name(name: str, used: set) -> str:
    """A valid, unique sheet name — Excel silently corrupts files without one."""
    cleaned = INVALID_SHEET_CHARS.sub("-", name).strip() or "Project"
    candidate = cleaned[:MAX_SHEET_NAME]
    suffix = 2
    while candidate.lower() in used:
        tail = f" ({suffix})"
        candidate = cleaned[:MAX_SHEET_NAME - len(tail)] + tail
        suffix += 1
    used.add(candidate.lower())
    return candidate


def export_xlsx(
    db: Database,
    projects: list[Project],
    date_from: Optional[str],
    date_to: Optional[str],
    worker: str,
    output_path: str,
    company: str = "",
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    def fill(color):
        return PatternFill("solid", fgColor=color)

    def font(bold=False, color=C_INK, size=10):
        return Font(bold=bold, color=color, size=size, name="Calibri")

    def align(h="left", v="center", wrap=False):
        return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

    hairline = Border(bottom=Side(border_style="thin", color="FFE3E8EF"))

    def widths(sheet, values):
        for index, width in enumerate(values, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    def header_row(sheet, row, labels, background=C_BAND):
        for column, label in enumerate(labels, 1):
            cell = sheet.cell(row, column, label)
            cell.font = font(bold=True, color=C_WHITE, size=9)
            cell.fill = fill(background)
            cell.alignment = align("left" if column == 1 else "center")
        sheet.row_dimensions[row].height = 20

    project_ids = [p.id for p in projects]
    wb = Workbook()

    # ── Summary sheet ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    title = ws.cell(row, 1, "WorkTrack — Timesheet Report")
    title.font = Font(bold=True, size=16, color=C_WHITE, name="Calibri")
    title.fill = fill(C_BAND)
    title.alignment = align("left")
    ws.row_dimensions[row].height = 32
    row += 2

    meta = [("Prepared by", worker or "—")]
    if company:
        meta.append(("Organisation", company))
    meta += [
        ("Reporting period", period_label(date_from, date_to)),
        ("Generated", now_local().strftime("%d %B %Y %H:%M")),
    ]
    for label, value in meta:
        key = ws.cell(row, 1, label)
        key.font = font(bold=True, color=C_MUTED, size=9)
        val = ws.cell(row, 2, value)
        val.font = font(size=10)
        row += 1
    row += 1

    header_row(ws, row, ["Project", "Client", "Task", "Sessions",
                         "Duration (H:MM)", "Hours"])
    header_start = row
    row += 1

    summary = db.get_sessions_summary(date_from, date_to, project_ids)
    grand_total = 0.0
    for index, entry in enumerate(summary):
        seconds = entry["total_seconds"] or 0.0
        grand_total += seconds
        values = [entry["project_name"], entry["client"] or "", entry["subtask_name"],
                  entry["session_count"], format_hm(seconds), round(seconds / 3600, 2)]
        for column, value in enumerate(values, 1):
            cell = ws.cell(row, column, value)
            cell.font = font(size=10)
            cell.alignment = align("left" if column <= 3 else "center")
            cell.border = hairline
            if index % 2:
                cell.fill = fill(C_ZEBRA)
            if column == 6:
                cell.number_format = "0.00"
        row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    total_label = ws.cell(row, 1, "GRAND TOTAL")
    total_label.font = font(bold=True, size=11)
    total_label.alignment = align("right")
    ws.cell(row, 5, format_hm(grand_total)).font = font(bold=True, size=11)
    total_hours = ws.cell(row, 6, round(grand_total / 3600, 2))
    total_hours.font = font(bold=True, color=C_ACCENT, size=12)
    total_hours.number_format = "0.00"
    for column in (5, 6):
        ws.cell(row, column).alignment = align("center")
    ws.row_dimensions[row].height = 22

    widths(ws, [32, 24, 40, 11, 16, 10])
    ws.freeze_panes = ws.cell(header_start + 1, 1)
    if row - header_start > 1:
        ws.auto_filter.ref = f"A{header_start}:F{row - 1}"

    # ── Per-project detail sheets ─────────────────────────────────────────
    used_names = {"summary"}
    for project in projects:
        sessions = db.get_sessions(project_id=project.id, date_from=date_from,
                                   date_to=date_to)
        if not sessions:
            continue
        sessions.sort(key=lambda s: parse_local(s.started_at) or now_local())

        sheet = wb.create_sheet(title=_safe_sheet_name(project.name, used_names))
        sheet.sheet_view.showGridLines = False

        r = 1
        sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        banner = sheet.cell(r, 1, project.name)
        banner.font = Font(bold=True, size=14, color=C_WHITE, name="Calibri")
        banner.fill = fill(C_BAND)
        banner.alignment = align("left")
        sheet.row_dimensions[r].height = 28
        r += 1

        if project.client:
            sheet.cell(r, 1, f"Client: {project.client}").font = font(
                color=C_MUTED, size=9)
            r += 1
        r += 1

        header_row(sheet, r, ["Date", "Start", "End", "Task",
                              "Duration (H:MM)", "Hours", "Notes"], C_HEAD)
        detail_header = r
        r += 1

        project_total = 0.0
        for index, s in enumerate(sessions):
            seconds = s.duration_seconds or 0.0
            project_total += seconds
            values = [
                fmt_date(s.started_at, "%Y-%m-%d"),
                fmt_time(s.started_at),
                fmt_time(s.ended_at),
                s.subtask_name,
                format_hm(seconds),
                round(seconds / 3600, 2),
                s.notes or "",
            ]
            for column, value in enumerate(values, 1):
                cell = sheet.cell(r, column, value)
                cell.font = font(size=10)
                cell.alignment = align("center" if column in (1, 2, 3, 5, 6) else "left")
                cell.border = hairline
                if index % 2:
                    cell.fill = fill(C_ZEBRA)
                if column == 6:
                    cell.number_format = "0.00"
            r += 1

        sheet.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        label = sheet.cell(r, 1, "PROJECT TOTAL")
        label.font = font(bold=True, size=10)
        label.alignment = align("right")
        sheet.cell(r, 5, format_hm(project_total)).font = font(bold=True, size=10)
        hours = sheet.cell(r, 6, round(project_total / 3600, 2))
        hours.font = font(bold=True, color=C_ACCENT, size=11)
        hours.number_format = "0.00"
        for column in (5, 6):
            sheet.cell(r, column).alignment = align("center")

        widths(sheet, [13, 9, 9, 38, 16, 10, 34])
        sheet.freeze_panes = sheet.cell(detail_header + 1, 1)

    wb.save(output_path)
